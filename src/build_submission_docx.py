from pathlib import Path
import re,subprocess
from docx import Document
from docx.shared import Pt, Mm, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
R=Path(__file__).resolve().parents[1]
d=Document();sec=d.sections[0];sec.page_width=Mm(210);sec.page_height=Mm(297)
sec.top_margin=Mm(17);sec.bottom_margin=Mm(17);sec.left_margin=Mm(18);sec.right_margin=Mm(18)
sec.header_distance=Mm(7);sec.footer_distance=Mm(8)
for name in ['Normal','Body Text','First Paragraph','Compact','Table','Caption']:
 try:s=d.styles[name]
 except KeyError:s=d.styles.add_style(name,1)
 s.font.name='Noto Sans CJK SC';s.font.size=Pt(9.5);s.font.color.rgb=RGBColor.from_string('19281F')
 s.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Noto Sans CJK SC')
 s.paragraph_format.line_spacing=1.2;s.paragraph_format.space_after=Pt(5)
for name,sz in [('Heading 1',18),('Heading 2',12),('Heading 3',11),('Heading 4',10)]:
 s=d.styles[name];s.font.name='Noto Sans CJK SC';s.font.size=Pt(sz);s.font.color.rgb=RGBColor.from_string('2A5E4E');s.font.bold=True
 s.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Noto Sans CJK SC')
 s.paragraph_format.space_before=Pt(9);s.paragraph_format.space_after=Pt(5);s.paragraph_format.keep_with_next=True
h=sec.header.paragraphs[0];h.text='VOICELOG / 声迹                                            预赛材料 · v1.1';h.runs[0].font.size=Pt(8);h.runs[0].font.color.rgb=RGBColor.from_string('76836F')
f=sec.footer.paragraphs[0];f.alignment=2;r=f.add_run('VoiceLog / 声迹 · ');r.font.size=Pt(8)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');f._p.append(field)
ref=R/'reports/docx_reference.docx';d.save(ref)
md=(R/'docs/VoiceLog_声迹_参赛材料_v1.1.md').read_text()
md=re.sub(r'```mermaid\n.*?```','![记录、理解与下一天的关系]('+str(R/'diagrams/参赛材料_01.png')+'){width=6.4in}',md,flags=re.S)
src=R/'reports/submission_docx_source.md';src.write_text(md)
out=R/'docs/VoiceLog_声迹_参赛材料_v1.1.docx'
subprocess.run(['pandoc',str(src),'-o',str(out),'--reference-doc',str(ref)],check=True)
d=Document(out)
for table in d.tables:
 table.autofit=False
 for row in table.rows:
  for cell in row.cells:
   for p in cell.paragraphs:
    p.paragraph_format.space_after=Pt(4);p.paragraph_format.space_before=Pt(4)
    for r in p.runs:r.font.size=Pt(9)
 for cell in table.rows[0].cells:
  shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'EDF2E4');cell._tc.get_or_add_tcPr().append(shade)
 trpr=table.rows[0]._tr.get_or_add_trPr();rh=OxmlElement('w:tblHeader');trpr.append(rh)
for p in d.paragraphs:
 for r in p.runs:
  if r._element.xpath('.//wp:docPr'):
   for el in r._element.xpath('.//wp:docPr'):el.set('descr','录音豆采集，经场景时间线与人物交汇，接到复盘和日程，延续至下一天。')
d.core_properties.author='VoiceLog';d.core_properties.title='VoiceLog / 声迹 · 预赛材料';d.save(out)
print(out)
